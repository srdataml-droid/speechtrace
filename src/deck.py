"""Extract text from a user-provided slide deck; no OCR or remote upload."""
from pathlib import Path


def slides_from_file(path: str | None, cap: int = 30) -> list[dict]:
    if not path:
        return []
    ext = Path(path).suffix.lower()
    if ext == ".pptx":
        from pptx import Presentation
        prs=Presentation(path)
        if len(prs.slides)>cap:
            raise ValueError(f"Deck exceeds {cap} slides.")
        result=[]
        for idx, slide in enumerate(prs.slides,1):
            paragraphs=[]
            for shape in slide.shapes:
                if getattr(shape,"has_text_frame",False):
                    t=shape.text.strip()
                    if t: paragraphs.append(t)
                elif getattr(shape,"has_table",False):
                    for row in shape.table.rows:
                        paragraphs.append(" ".join(cell.text for cell in row.cells))
            result.append({"number":idx,"text":"\n".join(paragraphs)[:4000]})
        return result
    if ext == ".pdf":
        from pypdf import PdfReader
        reader=PdfReader(path)
        if len(reader.pages)>cap:
            raise ValueError(f"PDF exceeds {cap} pages.")
        return [{"number":i, "text":(page.extract_text() or "")[:4000]}
                for i,page in enumerate(reader.pages,1)]
    if ext == ".txt":
        return [{"number":1,"text":Path(path).read_text(encoding="utf8")[:4000]}]
    raise ValueError("Supported deck formats: PPTX, PDF and TXT only.")
