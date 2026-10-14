from pathlib import Path
import numpy as np
import pandas as pd
import plotly.graph_objects as go
from pptx import Presentation
import pytest
from src.coaching import acoustic_evidence,transcript_evidence,deck_evidence,coach,SR
from src.deck import slides_from_file
from app import coach_rehearsal,demo


def synthetic_with_pause():
    t=np.arange(SR,dtype=np.float32)/SR
    tone=(.15*np.sin(2*np.pi*190*t)).astype(np.float32)
    return np.r_[tone,tone,np.zeros(int(1.25*SR),np.float32),tone,tone]


def test_pause_boundaries_and_report():
    e=acoustic_evidence(synthetic_with_pause())
    assert e['pause_count']>=1
    assert any(abs(x['start']-2)<.1 for x in e['events'] if x['type']=='low_energy_interval')
    assert e['duration_s']>5


def test_fillers_not_phantom_timestamps():
    e=transcript_evidence('Um we uh built a project. Uh we can improve.',12)
    assert e['filler_count']==3
    assert all(f['start'] is None for f in e['fillers'])
    assert e['wpm_total']>0


def test_asr_word_alignment_and_burst():
    word_seq=' '.join(['um']+['value']*21)
    events=[{'word':x,'start':i*.2,'end':i*.2+.15} for i,x in enumerate(word_seq.split())]
    result=transcript_evidence(word_seq,10,events)
    assert result['fillers'][0]['start']==0
    assert result['bursts']


def test_deck_overlap_not_factual_certainty():
    a=deck_evidence('We built a solution for water access', [{'number':1,'text':'Water access challenges'}])
    assert a['slides'][0]['matched_keywords']
    assert 'NOT semantic' in a['limitations'][0]


def test_pptx_text_extraction(tmp_path):
    prs=Presentation()
    slide=prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text='Improved water access'
    f=tmp_path/'deck.pptx';prs.save(f)
    assert 'water access' in slides_from_file(str(f))[0]['text']


def test_bad_deck_type(tmp_path):
    f=tmp_path/'bad.docx'; f.write_text('hey')
    with pytest.raises(ValueError,match='Supported deck'):
        slides_from_file(str(f))


def test_ui_callback_with_synthetic_audio():
    base=Path(__file__).resolve().parents[1]
    report=coach_rehearsal(str(base/'data/reference/ideal.wav'),None,'Hackathon judges',
                           'Manual transcript','Um I made a tool to help people learn.',False)
    assert 'Rehearsal evidence' in report[0]
    assert isinstance(report[2],go.Figure)
    assert isinstance(report[3],pd.DataFrame)
    assert report[5]['language']['filler_count']==1
    assert Path(report[6]).is_file()
    Path(report[6]).unlink()


def test_gradio_config():
    types={x['type'] for x in demo.config['components']}
    assert {'audio','file','plot','checkbox','dropdown'}.issubset(types)


def test_opt_in_gradio_report_can_be_read_by_local_mcp(tmp_path, monkeypatch):
    from service.mcp_server import read_report
    monkeypatch.setenv('SPEECHTRACE_REPORT_DB',str(tmp_path/'reports.sqlite'))
    base=Path(__file__).resolve().parents[1]
    summary, transcript, plot, events, review, preview, jsonpath = coach_rehearsal(
        str(base/'data/reference/ideal.wav'),None,'Pitch to judges',
        'Manual transcript','Um our idea helps everyone',False,False,True)
    assert 'Assistant report ID' in summary
    import json
    payload=json.loads(Path(jsonpath).read_text())
    report_id=payload['stored_report_id']
    assert report_id
    assert read_report(report_id)['language']['filler_count']==1
    Path(jsonpath).unlink()
